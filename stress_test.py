#!/usr/bin/env python3
"""
Stress tests: make the proxy data HARDER than the tuning data and see what breaks.
  A) BOM wording noise: typos, unseen abbreviations, reordered words (vector drawings)
  B) heavy scan degradation: lower DPI, more skew/blur/noise, strong JPEG (OCR path)
"""
import glob, os, random, re, shutil, sys, warnings
warnings.filterwarnings("ignore")
import openpyxl
sys.path.insert(0, os.path.join(os.path.dirname(os.path.abspath(__file__)), "samples"))
from evaluate import evaluate, HARD

ABBR = {"header": "hdr", "valve": "vlv", "flange": "flg", "plate": "plt", "economiser": "econ", "pressure": "press.", "temperature": "temp.",
        "inspection": "insp.", "expansion": "expn", "gasket": "gskt", "support": "supp."}

def mutate(text, rng):
    words = text.split(" ")
    for i, w in enumerate(words):
        lw = w.lower().strip(",")
        if lw in ABBR and rng.random() < 0.5:
            words[i] = w.lower().replace(lw, ABBR[lw])
        elif len(w) > 5 and w.isalpha() and rng.random() < 0.18:          # typo: swap two inner letters
            k = rng.randint(1, len(w) - 3); words[i] = w[:k] + w[k + 1] + w[k] + w[k + 2:]
    if len(words) > 3 and rng.random() < 0.3:
        words = words[1:] + words[:1]                                     # reorder: move first word to the end
    return " ".join(words)

def make_noisy_boms(src, dst, seed=1):
    import json
    rng = random.Random(seed)
    os.makedirs(dst, exist_ok=True)
    maps = {}
    for f in glob.glob(os.path.join(src, "sample_*")):
        if f.endswith("_bom.xlsx"):
            maps_k = os.path.basename(f).split("_")[1]
            maps[maps_k] = {}
            wb = openpyxl.load_workbook(f); ws = wb.active
            # find description column by header name
            hdr = next(r for r in range(1, 12) if any(str(c.value or "").lower() in ("description", "item description") for c in ws[r]))
            col = next(c.column for c in ws[hdr] if str(c.value or "").lower() in ("description", "item description"))
            for r in range(hdr + 1, ws.max_row + 1):
                v = ws.cell(r, col).value
                if v:
                    nv = mutate(str(v), rng); maps[maps_k][str(v)] = nv; ws.cell(r, col).value = nv
            wb.save(os.path.join(dst, os.path.basename(f)))
        elif not f.endswith("_truth.json"):
            shutil.copy(f, dst)
    for f in glob.glob(os.path.join(src, "sample_*_truth.json")):          # translate truth keys to the mutated wording
        t = json.load(open(f)); idx = os.path.basename(f).split("_")[1]
        for e in t["errors"]:
            if e["type"] in ("extra", "duplicate"):
                e["key"] = maps[idx].get(e["key"], e["key"])
        json.dump(t, open(os.path.join(dst, os.path.basename(f)), "w"))

def heavy_scans(src, dst, n=10):
    import fitz, numpy as np, cv2
    sys.path.insert(0, "samples")
    os.makedirs(dst, exist_ok=True)
    for f in glob.glob(os.path.join(src, "sample_*")):
        if not f.endswith("_drawing_scan.pdf"):
            shutil.copy(f, dst)
    pdfs = sorted(glob.glob(os.path.join(src, "sample_*_drawing.pdf")))[:n]
    for k, pdf in enumerate(pdfs):
        rng = np.random.default_rng(k)
        d = fitz.open(pdf); page = d[0]
        pix = page.get_pixmap(matrix=fitz.Matrix(150 / 72, 150 / 72), colorspace=fitz.csGRAY)            # only 150 dpi source
        img = np.frombuffer(pix.samples, dtype=np.uint8).reshape(pix.h, pix.w).copy()
        M = cv2.getRotationMatrix2D((pix.w / 2, pix.h / 2), rng.uniform(-0.9, 0.9), 1.0)                 # up to ~1 degree skew
        img = cv2.warpAffine(img, M, (pix.w, pix.h), borderValue=255)
        img = cv2.GaussianBlur(img, (5, 5), 1.4)
        img = np.clip(img + rng.normal(0, 14, img.shape), 0, 255).astype(np.uint8)
        # uneven illumination (shadow gradient)
        grad = np.tile(np.linspace(0.78, 1.0, img.shape[1]), (img.shape[0], 1)); img = (img * grad).astype(np.uint8)
        ok, enc = cv2.imencode(".jpg", img, [int(cv2.IMWRITE_JPEG_QUALITY), 40])
        out = fitz.open(); p = out.new_page(width=page.rect.width, height=page.rect.height); p.insert_image(p.rect, stream=enc.tobytes())
        out.save(os.path.join(dst, os.path.basename(pdf).replace("_drawing.pdf", "_drawing_scan.pdf")))

def report(title, res):
    stats, anyd, per, conf, n = res
    rel = stats.pop("_rel", None)
    T = sum(s["truth"] for s in stats.values()); TP = sum(s["tp"] for s in stats.values()); FP = sum(s["fp"] for s in stats.values())
    print(f"{title}\n  run reliability flagged: {rel}\n  samples {n} | injected errors {T} | found {TP} (recall {100*TP/T:.0f}%) | false alarms {FP} (precision {100*TP/max(1,TP+FP):.0f}%) | confirm flags/drawing {conf/n:.2f}")
    for t in HARD:
        s = stats[t]
        if s["truth"] and s["tp"] < s["truth"]: print(f"    weak: {t} {s['tp']}/{s['truth']}")
    bad = [p for p in per if p["missed"] or p["fps"]]
    for p in bad[:6]: print(f"    sample {p['sample']}: missed={p['missed']} false_alarms={p['fps']}")

if __name__ == "__main__":
    make_noisy_boms("samples_heldout", "samples_stress_a")
    report("A) NOISY BOM WORDING (typos / abbreviations / reordering), vector drawings", evaluate("samples_stress_a"))
    heavy_scans("samples_heldout", "samples_stress_b", n=10)
    report("B) HEAVY SCAN DEGRADATION (150 dpi, ~1 deg skew, blur, noise, shadow, JPEG q40), OCR path", evaluate("samples_stress_b", scans=True))
